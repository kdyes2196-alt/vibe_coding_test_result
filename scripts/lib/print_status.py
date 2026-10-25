"""생성된 현황 엑셀의 요약 시트를 기관별 표로 출력한다. 사용법: python3 print_status.py <xlsx> [이전 xlsx]"""
import sys
import openpyxl


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = {}
    for row in ws.iter_rows(min_row=6, max_row=40, min_col=3, max_col=19, values_only=True):
        if not row[0]:
            break
        m = [x or "-" for x in row[2:]]
        groups = ["".join(m[0:3]), m[3], m[4], m[5], "".join(m[6:12]), "".join(m[12:15])]
        evidence = []
        for r in wb[row[0]].iter_rows(min_row=5, values_only=True):
            if r[1] and r[6] and r[7]:
                evidence.append(f"{r[1]} {r[6]}/{r[7]}")
        rows[row[0]] = (row[1] or 0, groups, evidence)
    return rows


cur = read(sys.argv[1])
prev = read(sys.argv[2]) if len(sys.argv) > 2 else {}
print("  기관명        진척률(직전)  980001 980002 980003 980004  900N02  AC4111   원장 근거(확인일시/거래고유번호)")
for name, (progress, g, ev) in cur.items():
    delta = ""
    if name in prev:
        d = progress - prev[name][0]
        delta = f"({d:+.0f})" if d else "(  -)"
    mark = " ▲" if delta.startswith("(+") else ("  ▼" if delta.startswith("(-") and delta != "(  -)" else "")
    print(f"  {name:<10}{progress:>5.0f}% {delta:<6}  {g[0]:<6} {g[1]:<6} {g[2]:<6} {g[3]:<6}  {g[4]:<6}  {g[5]:<6}   {'; '.join(ev)}{mark}")
