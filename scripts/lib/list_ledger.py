"""원장 CSV(H2 CSVWRITE)를 표로 출력하고, 최근 생성한 현황 엑셀의 기관별 시트에 기재된 거래고유번호인지 함께 표시한다.
사용법: list_ledger.py <원장csv> <현황xlsx 또는 -> <샘플xlsx(기관명 매핑용)>"""
import csv
import sys
import openpyxl

ledger_csv, report, sample = sys.argv[1:4]

names = {}
ws = next(s for s in openpyxl.load_workbook(sample).worksheets if "(상세)" in s.title)
for r in ws.iter_rows(min_row=6, max_row=60, min_col=2, max_col=4, values_only=True):
    if not r[2]:
        break
    names[(str(r[0]).zfill(3), str(r[1]).zfill(3))] = r[2]

recorded = {}
if report != "-":
    wb = openpyxl.load_workbook(report, data_only=True)
    for agency in names.values():
        if agency not in wb.sheetnames:
            continue
        for r in wb[agency].iter_rows(min_row=5, values_only=True):
            if r[1] and r[7]:
                case = str(r[2]).split("\n")[0]
                recorded.setdefault(str(r[7]), []).append(f"{agency}/{r[1]} {case}")

rows = list(csv.DictReader(open(ledger_csv, encoding="utf-8")))
head = ["거래일시", "거래고유번호", "기관", "거래구분", "상태", "추가조회", "페이징", "재신규", "결과", "엑셀 기재(시트/테스트케이스)"]
table = []
for r in rows:
    agency = names.get((r["ACQ_INST_CD"], r["ISS_INST_CD"]), "?")
    rec = recorded.get(r["TRX_UNQ_NO"])
    if rec:
        note = "O " + "; ".join(rec)
    elif r["TRX_STAT_CD"] != "00":
        note = "- (비정상 거래)"
    else:
        note = "- (미기재)"
    table.append([r["TRX_DT"][:19], r["TRX_UNQ_NO"], f"{agency}({r['ACQ_INST_CD']}/{r['ISS_INST_CD']})",
                  r["TRX_DV_CD"], r["TRX_STAT_CD"], r["ADD_INQ_COND"] or "-", r["PAGING_YN"] or "-",
                  r["RENEW_YN"] or "-", r["RSLT_CD"] or "-", note])

width = lambda s: sum(2 if ord(c) > 0x1100 else 1 for c in s)
cols = [max(width(x) for x in [h] + [t[i] for t in table]) for i, h in enumerate(head)]
fmt = lambda row: " | ".join(c + " " * (cols[i] - width(c)) for i, c in enumerate(row))
print(fmt(head))
print("-+-".join("-" * c for c in cols))
for t in table:
    print(fmt(t))
print(f"({len(table)}건, 엑셀 기재 {sum(1 for t in table if t[-1].startswith('O'))}건"
      + (f" / 기준 엑셀: {report})" if report != "-" else " / 아직 실행 전이라 엑셀 없음)"))
